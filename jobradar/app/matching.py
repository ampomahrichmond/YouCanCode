"""Resume ↔ job matching, fully local, per user.

Text is tokenised into unigrams + within-phrase bigrams, weighted with TF-IDF
over the stored job corpus, and compared by cosine similarity. Two boosts sit
on top: how much of the job *title* the resume covers, and whether the job's
seniority bucket fits the seniority the resume reads at. Scores are 0–100 and
every score keeps a short "why" (shared skills / notable gaps)."""
from __future__ import annotations

import io
import json
import logging
import math
import re
from collections import Counter

from . import db

log = logging.getLogger("jobradar.matching")

# ── text extraction ───────────────────────────────────────────────────────
def extract_text(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    return data.decode("utf-8", errors="replace")


# ── tokenisation ──────────────────────────────────────────────────────────
STOP = set("""a an and are as at be been but by for from has have he her his i if in into is it its
me my no nor not of on or our she so than that the their them there these they this those to us
was we were what when where which while who whom why will with within without would you your yours
about above across after again all also am among any because before being below between both can
could did do does doing down during each few further here how more most other out over own same
should some such then through too under until up very via per etc ie eg re vs may might must
job jobs role roles position positions candidate candidates ideal preferred required requirements
responsibilities responsibility experience experiences years year work working works skills skill
ability able strong excellent team teams company companies opportunity opportunities including
include includes plus new well good great one two three four five well-known equal employer
applicants applicant application apply please email benefits salary pay compensation location
remote hybrid onsite full-time fulltime part-time time day days week weeks month months""".split())
_TOKEN_RX = re.compile(r"[a-z][a-z0-9+#.\-]{1,}")
_SEGMENT_RX = re.compile(r"[\n\r.,;:|•·()\[\]/]+|\s[-–—]\s")


def tokenize(text: str) -> list[str]:
    out: list[str] = []
    for seg in _SEGMENT_RX.split((text or "").lower()):
        words = [w.strip(".-") for w in _TOKEN_RX.findall(seg)]
        words = [w for w in words if len(w) >= 2 and w not in STOP and not w.isdigit()]
        out += words
        out += [f"{a} {b}" for a, b in zip(words, words[1:])]
    return out


# ── seniority read from a resume ──────────────────────────────────────────
_LEVEL_HINTS = {
    "C-level": r"\b(chief|cdo|cio|cto|cdao)\b", "VP": r"\b(vp|svp|evp|avp|vice[\s-]?president)\b",
    "Head": r"\bhead of\b", "Director": r"\bdirector\b", "Senior Manager": r"\b(senior|sr\.?)\s+manager\b",
    "Manager": r"\bmanager\b", "Lead": r"\b(lead|principal)\b", "Senior Analyst": r"\b(senior|sr\.?)\s+analyst\b",
    "Analyst": r"\banalyst\b",
}
LEVEL_RANK = {"Analyst": 1, "Senior Analyst": 2, "Lead": 3, "Manager": 3, "Senior Manager": 4,
              "Director": 5, "Head": 5, "VP": 6, "C-level": 7, "Other": 3}


def resume_levels(text: str) -> list[str]:
    counts = {k: len(re.findall(rx, text, re.I)) for k, rx in _LEVEL_HINTS.items()}
    return [k for k, c in sorted(counts.items(), key=lambda x: -x[1]) if c > 0]


def level_fit(resume_lvls: list[str], job_level: str | None) -> float:
    if not resume_lvls or not job_level or job_level == "Other":
        return 0.5
    d = LEVEL_RANK.get(job_level, 3) - LEVEL_RANK.get(resume_lvls[0], 3)
    return 1.0 if d in (0, 1) else 0.5 if d in (-1, 2) else 0.0


# ── TF-IDF model ──────────────────────────────────────────────────────────
class Model:
    def __init__(self, docs: dict[str, str]):
        self.n = max(1, len(docs))
        df: Counter = Counter()
        for text in docs.values():
            df.update(set(tokenize(text)))
        self.df = df

    def idf(self, term: str) -> float:
        return math.log((self.n + 1) / (self.df.get(term, 0) + 1)) + 1

    def vec(self, text: str) -> dict[str, float]:
        tf = Counter(tokenize(text))
        return {t: (1 + math.log(c)) * self.idf(t) for t, c in tf.items()}


def _cosine(a: dict[str, float], b: dict[str, float]) -> float:
    if not a or not b:
        return 0.0
    dot = sum(w * b[t] for t, w in a.items() if t in b)
    na = math.sqrt(sum(w * w for w in a.values()))
    nb = math.sqrt(sum(w * w for w in b.values()))
    return dot / (na * nb) if na and nb else 0.0


_COS_REF = 0.18


def score_job(model: Model, rvec, rtokens, rlevels, job: dict) -> tuple[int, dict]:
    text = f"{job.get('title') or ''}\n{job.get('description') or job.get('snippet') or ''}"
    jvec = model.vec(text)
    cos = _cosine(rvec, jvec)
    title_toks = [t for t in tokenize(job.get("title") or "") if " " not in t]
    title_cov = sum(1 for t in title_toks if t in rtokens) / len(title_toks) if title_toks else 0
    fit = level_fit(rlevels, job.get("level"))
    score = 100 * (0.65 * min(cos / _COS_REF, 1.0) + 0.25 * title_cov + 0.10 * fit)
    shared = sorted(((w * rvec[t], t) for t, w in jvec.items() if t in rvec), reverse=True)
    skills, seen = [], set()
    for _, t in shared:
        base = t.split(" ")[0]
        if base in seen or model.idf(t) < 1.4:
            continue
        seen.add(base); skills.append(t)
        if len(skills) == 8:
            break
    gaps = []
    for t, w in sorted(jvec.items(), key=lambda x: -x[1]):
        if t in rvec or " " in t or model.idf(t) < 2.0 or len(t) < 4:
            continue
        gaps.append(t)
        if len(gaps) == 5:
            break
    why = {"skills": skills, "gaps": gaps, "cos": round(cos, 3), "title": round(title_cov, 2), "fit": fit}
    return int(round(min(100, max(0, score)))), why


# ── per-user state ────────────────────────────────────────────────────────
_shared: dict = {"model": None, "n_docs": 0}
_users: dict[int, dict] = {}     # uid -> {rvec, rtokens, rlevels, resume_at}


def resume_text(uid: int) -> str | None:
    return db.get_setting("resume_text", uid)


def _model(jobs: list[dict] | None = None) -> Model:
    jobs = jobs if jobs is not None else db.jobs_for_scoring()
    if _shared["model"] is None or abs(len(jobs) - _shared["n_docs"]) > max(50, _shared["n_docs"] // 10):
        _shared["model"] = Model({j["id"]: f"{j['title']}\n{j.get('description') or j.get('snippet') or ''}" for j in jobs})
        _shared["n_docs"] = len(jobs)
    return _shared["model"]


def _user_state(uid: int) -> dict | None:
    resume = resume_text(uid)
    if not resume:
        _users.pop(uid, None)
        return None
    at = db.get_setting("resume_at", uid)
    st = _users.get(uid)
    if st and st["resume_at"] == at and st["model"] is _shared["model"]:
        return st
    model = _model()
    st = {"rvec": model.vec(resume), "rtokens": set(tokenize(resume)), "rlevels": resume_levels(resume),
          "resume_at": at, "model": model}
    _users[uid] = st
    return st


def rescore_all(uid: int) -> int:
    """Score every stored job against this user's resume (minutes of CPU on a
    40k-row corpus: run it in the sweep process, never in the web server).
    Also stores the resume summary the dashboard shows."""
    jobs = db.jobs_for_scoring()
    _shared["model"] = None
    model = _model(jobs)
    st = _user_state(uid)
    if not st:
        db.clear_scores(uid)
        db.set_setting("resume_summary", None, uid)
        return 0
    rows = []
    for j in jobs:
        s, why = score_job(model, st["rvec"], st["rtokens"], st["rlevels"], j)
        rows.append((s, json.dumps(why), j["id"]))
    db.update_scores(uid, rows)
    db.set_setting("rescore_n", str(len(jobs)), uid)
    db.set_setting("resume_summary", json.dumps(_summary(uid, st)), uid)
    log.info("rescored %d jobs for user %d", len(rows), uid)
    return len(rows)


def refresh(uid: int) -> int:
    """After a sweep: new rows were scored as they were stored, so a full
    rescore is only worth it when the corpus (and therefore the IDF weights)
    has grown by more than ~10 % since the last full rescore. The count is
    kept in settings so the decision survives across sweep processes."""
    if not resume_text(uid):
        return 0
    n = db.jobs_count()
    try:
        last = int(db.get_setting("rescore_n", uid) or 0)
    except ValueError:
        last = 0
    if last and abs(n - last) <= max(50, last // 10) and db.get_setting("resume_summary", uid):
        return 0
    return rescore_all(uid)


def score_batch(jobs: list[dict]) -> dict[int, list[tuple]]:
    """Score freshly scraped jobs for every user with a resume. Returns
    {uid: [(score, why_json, job_dict)]}; call db.update_scores after the upsert
    assigns ids."""
    out: dict[int, list[tuple]] = {}
    for uid in db.users_with_resume():
        st = _user_state(uid)
        if not st:
            continue
        out[uid] = [(*_score_pair(st, j), j) for j in jobs]
    return out


def _score_pair(st, j):
    s, why = score_job(st["model"], st["rvec"], st["rtokens"], st["rlevels"], j)
    return s, json.dumps(why)


def _summary(uid: int, st: dict) -> dict:
    rvec, model = st["rvec"], st["model"]
    top, seen = [], set()
    for t, _ in sorted(rvec.items(), key=lambda x: -x[1]):
        base = t.split(" ")[0]
        if base in seen or model.df.get(t, 0) < 2 or model.idf(t) < 1.3:
            continue
        seen.add(base); top.append(t)
        if len(top) == 20:
            break
    return {"present": True, "name": db.get_setting("resume_name", uid), "uploaded_at": st["resume_at"],
            "words": len((resume_text(uid) or "").split()), "levels": st["rlevels"][:3], "skills": top}


def resume_summary(uid: int) -> dict:
    """What the resume card shows. Read from the summary the last rescore
    stored, so the web process never has to build the corpus model (a minute
    of CPU). Before the first rescore finishes, a model-free stub is returned."""
    text = resume_text(uid)
    if not text:
        return {"present": False}
    raw = db.get_setting("resume_summary", uid)
    if raw:
        try:
            s = json.loads(raw)
            if s.get("uploaded_at") == db.get_setting("resume_at", uid):
                return s
        except json.JSONDecodeError:
            pass
    return {"present": True, "name": db.get_setting("resume_name", uid), "uploaded_at": db.get_setting("resume_at", uid),
            "words": len(text.split()), "levels": resume_levels(text)[:3], "skills": [], "scoring": True}
