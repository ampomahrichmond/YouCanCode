"""Markdown → .docx for the tailored resume and cover letter, so the draft can
be edited in Word and attached to an application."""
from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

_INLINE_RX = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|_[^_]+_|`[^`]+`)")
_BULLET_RX = re.compile(r"^\s*([-*•]|\d+[.)])\s+(.*)$")
_HEAD_RX = re.compile(r"^(#{1,6})\s+(.*)$")


def _add_runs(par, text: str, base_bold=False):
    """Split **bold**, *italic*, _italic_, `code` into runs."""
    for chunk in _INLINE_RX.split(text):
        if not chunk:
            continue
        run = None
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            run = par.add_run(chunk[2:-2]); run.bold = True
        elif (chunk.startswith("*") and chunk.endswith("*") or chunk.startswith("_") and chunk.endswith("_")) and len(chunk) > 2:
            run = par.add_run(chunk[1:-1]); run.italic = True
        elif chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            run = par.add_run(chunk[1:-1]); run.font.name = "Consolas"
        else:
            run = par.add_run(chunk)
        if base_bold:
            run.bold = True
    return par


def _style(doc: Document):
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.space_before = Pt(0)
    for name, size, color in (("Heading 1", 16, "1B5240"), ("Heading 2", 12.5, "1B5240"), ("Heading 3", 11, "26333A")):
        h = doc.styles[name]
        h.font.name = "Calibri"
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = RGBColor.from_string(color)
        h.paragraph_format.space_before = Pt(8 if name != "Heading 1" else 0)
        h.paragraph_format.space_after = Pt(2)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Pt(50)
        s.left_margin = s.right_margin = Pt(58)


def md_to_docx(md: str, title_centered: bool = True) -> bytes:
    doc = Document()
    _style(doc)
    lines = (md or "").replace("\r\n", "\n").split("\n")
    i, first_heading_done = 0, False
    while i < len(lines):
        line = lines[i].rstrip()
        i += 1
        if not line.strip():
            continue
        if re.fullmatch(r"\s*(-{3,}|\*{3,}|_{3,})\s*", line):
            continue                       # horizontal rule
        m = _HEAD_RX.match(line)
        if m:
            level = min(3, len(m.group(1)))
            text = m.group(2).strip().strip("#").strip()
            if level == 1 and not first_heading_done:
                p = doc.add_heading("", level=1)
                _add_runs(p, text)
                if title_centered:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                first_heading_done = True
                # a contact line directly under the name is centered too
                if i < len(lines) and lines[i].strip() and not _HEAD_RX.match(lines[i]) and not _BULLET_RX.match(lines[i]):
                    c = doc.add_paragraph()
                    _add_runs(c, lines[i].strip())
                    if title_centered:
                        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    i += 1
            else:
                p = doc.add_heading("", level=level)
                _add_runs(p, text)
            continue
        m = _BULLET_RX.match(line)
        if m:
            style = "List Number" if m.group(1)[0].isdigit() else "List Bullet"
            p = doc.add_paragraph(style=style)
            _add_runs(p, m.group(2).strip())
            p.paragraph_format.space_after = Pt(1)
            continue
        # a whole-line bold label ("**Summary**") acts as a small heading
        if re.fullmatch(r"\*\*[^*]+\*\*:?", line.strip()):
            p = doc.add_heading("", level=3)
            _add_runs(p, line.strip().strip(":").strip("*"))
            continue
        # merge soft-wrapped lines of one paragraph
        para = [line.strip()]
        while i < len(lines) and lines[i].strip() and not _HEAD_RX.match(lines[i]) and not _BULLET_RX.match(lines[i]) \
                and not re.fullmatch(r"\*\*[^*]+\*\*:?", lines[i].strip()):
            para.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        _add_runs(p, " ".join(para))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
